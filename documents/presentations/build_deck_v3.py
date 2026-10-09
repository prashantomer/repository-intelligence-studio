"""
Build: Repository Intelligence Studio — presentation v3
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(0x14, 0x24, 0x3D)
TEAL = RGBColor(0x0D, 0x9E, 0x8B)
TEAL_MID = RGBColor(0x0B, 0x7A, 0x6C)
SLATE = RGBColor(0x4B, 0x5E, 0x7A)
OFFWHITE = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BG = RGBColor(0xEB, 0xF1, 0xF8)
DIVIDER = RGBColor(0xD0, 0xDC, 0xEC)
HIGHLIGHT = RGBColor(0xE8, 0xF5, 0xF2)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


def add_rect(slide, left, top, width, height, fill=None, line=None, line_w=Pt(0)):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.line.width = line_w
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, text, left, top, width, height, size=Pt(14), bold=False,
             color=NAVY, align=PP_ALIGN.LEFT, wrap=True, italic=False):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txb


def slide_bg(slide, color=OFFWHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_bar(slide, color=TEAL, top=Inches(0), height=Inches(0.06)):
    add_rect(slide, 0, top, W, height, fill=color)


def add_label_chip(slide, text, left, top, width=Inches(2.2), height=Inches(0.36),
                   bg=TEAL, fg=WHITE, size=Pt(10)):
    add_rect(slide, left, top, width, height, fill=bg)
    txb = slide.shapes.add_textbox(left, top + Inches(0.02), width, height)
    p = txb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text.upper()
    run.font.size = size
    run.font.bold = True
    run.font.color.rgb = fg


def slide_cover():
    slide = prs.slides.add_slide(BLANK)
    slide_bg(slide, NAVY)
    add_rect(slide, 0, 0, Inches(0.55), H, fill=TEAL)

    circ = slide.shapes.add_shape(9, W - Inches(2.8), -Inches(0.6), Inches(3.4), Inches(3.4))
    circ.fill.solid()
    circ.fill.fore_color.rgb = RGBColor(0x1C, 0x36, 0x5C)
    circ.line.fill.background()

    add_label_chip(slide, "Capstone MVP  ·  Feature Walkthrough",
                   Inches(0.85), Inches(1.6), width=Inches(4.2), bg=TEAL_MID, size=Pt(10))
    add_text(slide, "Repository\nIntelligence Studio",
             Inches(0.85), Inches(2.05), Inches(9), Inches(2.4),
             size=Pt(52), bold=True, color=WHITE)
    add_text(slide,
             "Ingestion · search · repository Q&A · impact analysis · deletion auditability",
             Inches(0.85), Inches(4.55), Inches(10.2), Inches(0.5),
             size=Pt(17), color=RGBColor(0xA8, 0xC4, 0xDC))
    add_rect(slide, Inches(0.85), Inches(6.75), Inches(11.6), Inches(0.03), fill=TEAL_MID)
    add_text(slide, "July 2026",
             Inches(0.85), Inches(6.85), Inches(3), Inches(0.4),
             size=Pt(11), color=RGBColor(0x70, 0x90, 0xB0))


FEATURES = [
    ("01", "Repository Creation & Ingestion",
     "Register a repository, ingest the tracked branch, and build a reusable indexed snapshot."),
    ("02", "Manual Re-Sync",
     "Force-refresh the repository snapshot, embeddings, and graph when the branch or settings change."),
    ("03", "Repository Assistant",
     "Ask repository questions with async, thread-aware, evidence-grounded responses."),
    ("04", "Semantic Search",
     "Find relevant code using vector retrieval plus exact symbol hits and snippet-style previews."),
    ("05", "Impact Analyzer",
     "Estimate blast radius with dependency traversal, saved reports, and narrated guidance."),
    ("06", "Provider Logs",
     "Inspect every assistant, embedding, and narration call with usage, latency, and status."),
    ("07", "Centralized AI Settings",
     "Control providers and models once at the user level for assistant and retrieval behavior."),
    ("08", "Hard Delete & Deletion Logs",
     "Permanently remove a repository and keep one surviving operational deletion snapshot."),
]


def slide_overview():
    slide = prs.slides.add_slide(BLANK)
    slide_bg(slide, OFFWHITE)
    add_accent_bar(slide, TEAL)
    add_label_chip(slide, "Feature Overview", Inches(0.5), Inches(0.18),
                   width=Inches(2.0), height=Inches(0.3), bg=TEAL, size=Pt(9))
    add_text(slide, "What The Current Build Delivers",
             Inches(0.5), Inches(0.55), Inches(12), Inches(0.6),
             size=Pt(28), bold=True, color=NAVY)
    add_text(slide,
             "Eight implemented product features covering repository ingestion, retrieval, reasoning, operations, and cleanup.",
             Inches(0.5), Inches(1.1), Inches(12.2), Inches(0.4),
             size=Pt(13), color=SLATE)

    cols = 2
    card_w = Inches(6.0)
    card_h = Inches(0.76)
    start_y = Inches(1.65)
    gap_x = Inches(0.5)
    gap_y = Inches(0.1)
    pad_x = Inches(0.5)

    for idx, (num, title, desc) in enumerate(FEATURES):
        col = idx % cols
        row = idx // cols
        left = pad_x + col * (card_w + gap_x)
        top = start_y + row * (card_h + gap_y)
        add_rect(slide, left, top, card_w, card_h, fill=WHITE, line=DIVIDER, line_w=Pt(0.8))
        add_rect(slide, left, top, Inches(0.55), card_h, fill=TEAL)
        add_text(slide, num, left, top + Inches(0.17), Inches(0.55), Inches(0.4),
                 size=Pt(13), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(slide, title, left + Inches(0.65), top + Inches(0.06), card_w - Inches(0.75), Inches(0.32),
                 size=Pt(13), bold=True, color=NAVY)
        add_text(slide, desc, left + Inches(0.65), top + Inches(0.38), card_w - Inches(0.75), Inches(0.34),
                 size=Pt(10.3), color=SLATE)


def feature_slide(num, title, subtitle, inputs, processing, outputs, flow, use_cases):
    slide = prs.slides.add_slide(BLANK)
    slide_bg(slide, OFFWHITE)
    add_accent_bar(slide, TEAL)
    add_label_chip(slide, f"Feature {num}", Inches(0.5), Inches(0.18),
                   width=Inches(1.55), height=Inches(0.3), bg=TEAL, size=Pt(9))
    add_text(slide, title, Inches(0.5), Inches(0.55), Inches(12.3), Inches(0.55),
             size=Pt(26), bold=True, color=NAVY)
    add_text(slide, subtitle, Inches(0.5), Inches(1.08), Inches(12), Inches(0.38),
             size=Pt(12.5), color=SLATE, italic=True)
    add_rect(slide, Inches(0.5), Inches(1.48), Inches(12.3), Inches(0.025), fill=DIVIDER)

    card_w = Inches(3.85)
    card_h = Inches(4.45)
    start_x = Inches(0.5)
    gap = Inches(0.28)
    top = Inches(1.6)

    sections = [
        ("Input", inputs, CARD_BG, TEAL),
        ("Processing", processing, HIGHLIGHT, TEAL_MID),
        ("Output", outputs, CARD_BG, TEAL),
    ]

    for index, (label, bullets, bg, accent) in enumerate(sections):
        left = start_x + index * (card_w + gap)
        add_rect(slide, left, top, card_w, card_h, fill=bg, line=DIVIDER, line_w=Pt(0.6))
        add_rect(slide, left, top, card_w, Inches(0.05), fill=accent)
        add_text(slide, label, left + Inches(0.2), top + Inches(0.12), card_w - Inches(0.3), Inches(0.36),
                 size=Pt(12), bold=True, color=accent)

        y_off = top + Inches(0.52)
        for bullet in bullets:
            txb = slide.shapes.add_textbox(left + Inches(0.2), y_off, card_w - Inches(0.35), Inches(0.4))
            tf = txb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = f"›  {bullet}"
            run.font.size = Pt(11.2)
            run.font.color.rgb = SLATE
            y_off += Inches(0.46)

    use_top = Inches(6.08)
    add_rect(slide, Inches(0.5), use_top, Inches(12.3), Inches(0.48), fill=WHITE, line=DIVIDER, line_w=Pt(0.6))
    add_text(slide, "Practical use", Inches(0.68), use_top + Inches(0.09), Inches(1.3), Inches(0.2),
             size=Pt(10), bold=True, color=TEAL_MID)
    add_text(slide, " • ".join(use_cases), Inches(1.75), use_top + Inches(0.08), Inches(10.8), Inches(0.24),
             size=Pt(10), color=SLATE)

    flow_top = Inches(6.6)
    add_rect(slide, Inches(0.5), flow_top, Inches(12.3), Inches(0.45), fill=NAVY)
    add_text(slide, "  " + flow, Inches(0.55), flow_top + Inches(0.09), Inches(12.1), Inches(0.28),
             size=Pt(10.2), color=RGBColor(0xA8, 0xD0, 0xC8))
    add_text(slide, str(int(num) + 2), W - Inches(0.65), H - Inches(0.4), Inches(0.5), Inches(0.35),
             size=Pt(10), color=SLATE, align=PP_ALIGN.CENTER)


SLIDES = [
    dict(
        num="01",
        title="Repository Creation & Initial Ingestion",
        subtitle="Register a public repository and build the first repository-scoped indexed snapshot.",
        inputs=[
            "Repository name and source URL",
            "Default branch or tracked-branch override",
            "User triggers initial indexing",
        ],
        processing=[
            "Create repository and ingestion records",
            "Clone source into disposable workspace",
            "Inventory files, parse entities, chunk content",
            "Generate embeddings and dependency edges",
            "Persist repository-scoped snapshot state",
        ],
        outputs=[
            "Indexed files, chunks, entities, routes, edges",
            "Tracked branch + commit SHA snapshot",
            "Ready state for search, assistant, and impact",
        ],
        flow="User → Repository Create → StartService → IngestionJob → Clone → Parse → Chunk → Embed → Graph Build",
        use_cases=[
            "Onboard a new repository once and reuse the indexed snapshot across all later workflows",
            "Reduce repeated manual codebase reading before asking architecture or change questions",
            "Cost saving: one ingestion snapshot supports many searches and assistant queries instead of repeated raw analysis",
        ],
    ),
    dict(
        num="02",
        title="Manual Re-Sync",
        subtitle="Refresh the repository snapshot on demand when code or embedding assumptions change.",
        inputs=[
            "Repository selected",
            "User clicks Re-sync",
            "Current tracked branch snapshot",
        ],
        processing=[
            "Start forced ingestion even after prior indexing",
            "Rebuild entities, chunks, embeddings, and edges",
            "Record ingestion history and audit events",
            "Clean temporary cloned workspace afterward",
        ],
        outputs=[
            "Fresh indexed snapshot",
            "Updated ingestion history and latest status",
            "Safe recovery after embedding/config changes",
        ],
        flow="User → Re-sync Action → StartService(force=true) → IngestionJob → Rebuilt Repository Snapshot",
        use_cases=[
            "Refresh the knowledge base after branch updates without re-adding the repository",
            "Rebuild vectors safely after embedding-model changes",
            "Cost saving: avoids stale-answer debugging and keeps retrieval aligned to the latest indexed branch state",
        ],
    ),
    dict(
        num="03",
        title="Repository Assistant",
        subtitle="Async, repository-scoped Q&A with follow-up context, structured answers, and grounded citations.",
        inputs=[
            "Natural-language repository question",
            "Active conversation thread",
            "Indexed repository snapshot",
        ],
        processing=[
            "Queue user question and create pending assistant message",
            "Reuse recent thread context for follow-up resolution",
            "Try structured inventory or flow-aware answer path first",
            "Fallback to retrieval + provider/local generation",
            "Stream finalized response back into the chat UI",
        ],
        outputs=[
            "Grounded repository answer",
            "Citations with file paths and line ranges",
            "Persistent conversation + provider usage log",
        ],
        flow="User → QueueQuestionService → AssistantResponseJob → Context Build → Structured/Flow Path or Retrieval → Answer → Turbo Update",
        use_cases=[
            "Ask architecture, dependency, and execution-flow questions without opening many files manually",
            "Use follow-up questions in the same thread to continue investigation faster",
            "Cost saving: structured answer paths avoid unnecessary full-provider generation for narrow inventory questions",
        ],
    ),
    dict(
        num="04",
        title="Semantic Search",
        subtitle="Repository code search backed by embeddings, exact symbol hits, and snippet-style previews.",
        inputs=[
            "Search query, concept, or symbol name",
            "Repository-scoped indexed chunks",
        ],
        processing=[
            "Embed query using active embedding profile",
            "Rank vector-nearest chunks inside the repository",
            "Boost exact in-chunk line hits for concrete symbols",
            "Build contextual snippet previews around matched lines",
            "Refresh results in place via Turbo",
        ],
        outputs=[
            "Ranked matching chunks",
            "Snippet previews with surrounding lines",
            "File path, type, and line-window evidence",
        ],
        flow="User → Search Query → Query Embedding → Vector Ranking + Exact Hit Scoring → Snippet Build → Turbo Results",
        use_cases=[
            "Find a method, controller, or workflow even when the exact file is unknown",
            "Use symbol search plus semantic ranking for faster code discovery during debugging or onboarding",
            "Cost saving: direct search resolves many questions before invoking the assistant pipeline",
        ],
    ),
    dict(
        num="05",
        title="Impact Analyzer",
        subtitle="Change-risk analysis using target resolution, graph traversal, saved reports, and narrated guidance.",
        inputs=[
            "Entity identifier or natural-language impact question",
            "Repository dependency graph and indexed evidence",
        ],
        processing=[
            "Interpret and resolve the likely code target",
            "Traverse upstream and downstream relationships",
            "Collect related files, jobs, routes, and evidence chunks",
            "Assign deterministic risk level",
            "Add provider-backed narration with safe fallback and save report",
        ],
        outputs=[
            "Impact summary with risk level",
            "Saved report with drill-down evidence",
            "Review checklist and operational guidance",
        ],
        flow="User → Query Interpreter → Entity Resolve → Graph Traversal → Evidence Retrieval → Impact Report → Saved History",
        use_cases=[
            "Check blast radius before refactoring a controller, service, model, or job",
            "Use saved reports during review discussions to compare risky change areas",
            "Cost saving: catches risky dependencies earlier and reduces rework from avoidable regressions",
        ],
    ),
    dict(
        num="06",
        title="Provider Logs",
        subtitle="Operational visibility into assistant, embedding, and narration calls across repository workflows.",
        inputs=[
            "Assistant calls",
            "Embedding calls",
            "Impact narration/provider requests",
        ],
        processing=[
            "Record provider, model, operation, and endpoint",
            "Capture status, latency, token usage, and metadata",
            "Estimate cost where supported",
            "Expose filterable UI history for review",
        ],
        outputs=[
            "Call history by provider and operation",
            "Debuggable failed-call trail",
            "Usage and cost auditability",
        ],
        flow="Feature Runtime Call → ProviderCallLogRecorder → Database Entry → Logs Workspace",
        use_cases=[
            "Trace why an assistant or embedding request failed without reading raw server logs",
            "Compare latency and token usage across providers/models",
            "Cost saving: exposes expensive or noisy provider usage so runtime configuration can be optimized",
        ],
    ),
    dict(
        num="07",
        title="Centralized AI Settings",
        subtitle="User-level control of assistant and embedding providers/models across all repositories.",
        inputs=[
            "Assistant provider and model",
            "Embedding provider and model",
            "Ollama/local endpoint or remote provider config",
        ],
        processing=[
            "Persist user-level runtime profile",
            "Update provider/model choices dynamically in UI",
            "Apply inherited settings across repository workspaces",
            "Surface re-index implications when embeddings change",
        ],
        outputs=[
            "Single source of truth for AI runtime behavior",
            "Consistent assistant, search, and impact configuration",
            "Clear provider capability boundaries",
        ],
        flow="User → Settings Screen → Save Profile → Repository Features Inherit Runtime Configuration",
        use_cases=[
            "Switch providers/models once and apply the profile across all repositories",
            "Keep assistant and embedding behavior consistent during demos or experiments",
            "Cost saving: easier provider/model tuning helps move workloads to cheaper or local options when acceptable",
        ],
    ),
    dict(
        num="08",
        title="Hard Delete & Deletion Logs",
        subtitle="Permanent repository removal with full cleanup and one surviving operational deletion snapshot.",
        inputs=[
            "Repository delete request with confirmation",
            "Repository-scoped records and temp workspace path",
        ],
        processing=[
            "Snapshot counts and repository metadata for deletion log",
            "Hard-delete repository-owned records and histories",
            "Remove cloned workspaces from disk",
            "Preserve one deletion-log row per deleted repository id",
        ],
        outputs=[
            "Repository fully removed from active system state",
            "Read-only deletion history for operational review",
            "Re-addition remains possible as a new repository lifecycle",
        ],
        flow="User → DestroyService → Snapshot Counts → Hard Delete Records → Workspace Cleanup → Deletion Log",
        use_cases=[
            "Remove obsolete or test repositories cleanly when they are no longer needed",
            "Keep a minimal operational record of what was deleted and cleaned",
            "Cost saving: reduces storage, indexed data volume, and future background-processing footprint",
        ],
    ),
]


def slide_use_cases():
    slide = prs.slides.add_slide(BLANK)
    slide_bg(slide, OFFWHITE)
    add_accent_bar(slide, TEAL)
    add_label_chip(slide, "Practical Use Cases", Inches(0.5), Inches(0.18),
                   width=Inches(2.3), height=Inches(0.3), bg=TEAL, size=Pt(9))
    add_text(slide, "Where This MVP Is Useful In Practice",
             Inches(0.5), Inches(0.55), Inches(12), Inches(0.6),
             size=Pt(28), bold=True, color=NAVY)
    add_text(slide,
             "Representative day-to-day engineering scenarios enabled by the current repository intelligence workflow.",
             Inches(0.5), Inches(1.08), Inches(12.0), Inches(0.38),
             size=Pt(12.5), color=SLATE, italic=True)
    add_rect(slide, Inches(0.5), Inches(1.48), Inches(12.3), Inches(0.025), fill=DIVIDER)

    use_cases = [
        ("New engineer onboarding",
         "Ask for architecture summaries, service locations, execution flow, and main integration points without manual file hunting."),
        ("Safe change planning",
         "Run Impact Analyzer before changing a controller, service, model, or job to estimate blast radius and review checkpoints."),
        ("Fast code discovery",
         "Use semantic search to find symbols, workflows, or concepts even when exact filenames or modules are not already known."),
        ("Operational debugging",
         "Inspect provider logs to verify which model was called, how long it took, whether it failed, and what usage it consumed."),
        ("Repository refresh after changes",
         "Re-sync a tracked branch after code changes or embedding-configuration changes so search and assistant answers stay consistent."),
        ("Repository retirement / cleanup",
         "Hard-delete a repository when it is no longer needed while still retaining one operational deletion snapshot for auditability."),
    ]

    card_left = Inches(0.65)
    card_top = Inches(1.8)
    card_w = Inches(12.0)
    card_h = Inches(0.72)
    gap = Inches(0.12)

    for index, (title, desc) in enumerate(use_cases):
        top = card_top + index * (card_h + gap)
        add_rect(slide, card_left, top, card_w, card_h, fill=WHITE, line=DIVIDER, line_w=Pt(0.8))
        add_rect(slide, card_left, top, Inches(0.12), card_h, fill=TEAL)
        add_text(slide, title, card_left + Inches(0.22), top + Inches(0.09), Inches(3.3), Inches(0.24),
                 size=Pt(12.5), bold=True, color=NAVY)
        add_text(slide, desc, card_left + Inches(3.1), top + Inches(0.09), Inches(8.5), Inches(0.42),
                 size=Pt(10.8), color=SLATE)

    add_text(slide, "11", W - Inches(0.65), H - Inches(0.4), Inches(0.5), Inches(0.35),
             size=Pt(10), color=SLATE, align=PP_ALIGN.CENTER)


def slide_closing():
    slide = prs.slides.add_slide(BLANK)
    slide_bg(slide, NAVY)
    add_rect(slide, 0, 0, Inches(0.55), H, fill=TEAL)
    circ = slide.shapes.add_shape(9, W - Inches(3.5), H - Inches(2.5), Inches(4.2), Inches(4.2))
    circ.fill.solid()
    circ.fill.fore_color.rgb = RGBColor(0x1C, 0x36, 0x5C)
    circ.line.fill.background()
    add_text(slide, "Thank You", Inches(0.85), Inches(2.2), Inches(10), Inches(1.4),
             size=Pt(52), bold=True, color=WHITE)
    add_text(slide, "Repository Intelligence Studio  ·  MVP Feature Walkthrough  ·  July 2026",
             Inches(0.85), Inches(3.75), Inches(10), Inches(0.45),
             size=Pt(14), color=RGBColor(0xA8, 0xC4, 0xDC))
    add_rect(slide, Inches(0.85), Inches(4.35), Inches(4), Inches(0.04), fill=TEAL_MID)


slide_cover()
slide_overview()
for slide_data in SLIDES:
    feature_slide(**slide_data)
slide_use_cases()
slide_closing()

output_path = Path(__file__).with_name("features-v3.pptx")
prs.save(output_path)
print(f"Saved: {output_path}")
